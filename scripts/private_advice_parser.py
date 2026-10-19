"""
Parse Borpathar private-sale GST tax invoices (.xlsx), one invoice per file.

These are the NON-auction sales (HUL, Tata, direct buyers). Each invoice carries
grade-wise lines: GRADE, Quantity (kg), Rate (INR/kg) and Amount (taxable, pre-GST).
We use the taxable line amount for realization — GST is a pass-through, not revenue.

Layout is label-anchored (columns found by their header text) so it survives
invoices with a different number of lines. Real buyer = Ship-to (HUL/Tata);
Bill-to is often a broker/buying arm (e.g. Select Commodities for HUL).
"""
import re
import openpyxl

_MON = {'JAN':1,'FEB':2,'MAR':3,'APR':4,'MAY':5,'JUN':6,
        'JUL':7,'AUG':8,'SEP':9,'OCT':10,'NOV':11,'DEC':12}


def _date(v):
    if v is None:
        return None
    if hasattr(v, "date"):
        try: return v.date().isoformat()
        except Exception: pass
    m = re.search(r"(\d{1,2})-([A-Za-z]{3})-(\d{2,4})", str(v))
    if m:
        d = int(m.group(1)); mo = _MON.get(m.group(2).upper())
        y = int(m.group(3)); y = y + 2000 if y < 100 else y
        if mo: return f"{y:04d}-{mo:02d}-{d:02d}"
    return None


def parse_private_invoice(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    maxr = min(ws.max_row, 100); maxc = min(ws.max_column, 24)

    cols = {}; grade_row = None
    inv_no = date = bill_to = ship_to = None
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            s = v.strip()
            if s == "Quantity": cols["qty"] = c
            elif s == "Rate": cols["rate"] = c
            elif s == "Amount": cols["amt"] = c
            elif s == "Sl": cols["sl"] = c
            elif s == "GRADE": cols["grade"] = c; grade_row = r
            elif s == "Invoice No.": inv_no = ws.cell(r + 1, c).value
            elif s == "Invoice Date": date = _date(ws.cell(r + 1, c).value)
            elif s.startswith("Buyer"): bill_to = ws.cell(r + 1, c).value
            elif "Ship to" in s or s.startswith("Delivery Address"):
                ship_to = ws.cell(r + 1, c).value

    sl = cols.get("sl", 1); gcol = cols.get("grade", 5)
    qcol = cols.get("qty", 6); rcol = cols.get("rate", 8); acol = cols.get("amt", 11)
    start = (grade_row or 22) + 1

    lines = []
    for r in range(start, maxr + 1):
        a = ws.cell(r, sl).value
        g = ws.cell(r, gcol).value
        q = ws.cell(r, qcol).value
        if isinstance(a, str) and a.strip().lower() == "total":
            break
        if isinstance(a, (int, float)) and g and isinstance(q, (int, float)):
            rate = ws.cell(r, rcol).value
            amt = ws.cell(r, acol).value
            lines.append({
                "grade": str(g).strip().upper(),
                "qty": float(q),
                "rate": float(rate) if isinstance(rate, (int, float)) else None,
                "amount": float(amt) if isinstance(amt, (int, float)) else None,
            })

    buyer = (ship_to or bill_to or "").strip() or None
    return {
        "path": path,
        "invoice": str(inv_no).strip() if inv_no else None,
        "date": date,
        "buyer": buyer,
        "bill_to": (bill_to or "").strip() or None,
        "lines": lines,
        "sold_kg": round(sum(l["qty"] for l in lines), 3),
        "sold_val": round(sum(l["amount"] or 0 for l in lines), 2),
    }


if __name__ == "__main__":
    import sys, glob, os
    paths = []
    for arg in sys.argv[1:]:
        paths += glob.glob(arg) if os.path.isfile(arg) else glob.glob(os.path.join(arg, "**", "*.xls*"), recursive=True)
    tk = tv = 0.0
    for p in paths:
        inv = parse_private_invoice(p)
        print(f"{os.path.basename(p)}: {inv['invoice']} {inv['date']} -> {inv['buyer']}")
        for l in inv["lines"]:
            print(f"    {l['grade']:<8} {l['qty']:>9.1f} kg @ {l['rate']}  = {l['amount']}")
        print(f"    sold_kg={inv['sold_kg']}  sold_val={inv['sold_val']:,.2f}")
        tk += inv["sold_kg"]; tv += inv["sold_val"]
    print(f"\nTOTAL private sold_kg={tk:,.1f}  sold_val={tv:,.2f}")
