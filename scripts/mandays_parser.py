"""Reads a BTC fortnightly wages workbook -> that period's TOTAL mandays.
Reads the 'Mandays' sheet's TOTAL cell exactly as the sheet computes it
(anchored on the 'Man Days' header column + the 'TOTAL' row, so it survives
row shifts). Returns {period, start, end, total_mandays}."""
import openpyxl, re
from datetime import date

def _dates(s):
    ds=re.findall(r'(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})', s or '')
    out=[]
    for d,m,y in ds:
        y=int(y); y=y+2000 if y<100 else y
        try: out.append(date(y,int(m),int(d)).isoformat())
        except: pass
    return (out[0] if out else None, out[1] if len(out)>1 else None)

def parse_mandays(path):
    wb=openpyxl.load_workbook(path, data_only=True)
    if 'Mandays' not in wb.sheetnames: return None
    ws=wb['Mandays']
    g=[[c.value for c in row] for row in ws.iter_rows()]
    period=None
    for row in g:
        for v in row:
            if isinstance(v,str) and 'period from' in v.lower():
                period=' '.join(v.split()); break
        if period: break
    # Man Days header column
    mcol=None
    for row in g:
        for ci,v in enumerate(row):
            if isinstance(v,str) and 'man day' in v.lower().replace('.','').replace("'",''):
                mcol=ci
    # TOTAL row -> value in mcol
    total=None
    for row in g:
        for ci,v in enumerate(row):
            if isinstance(v,str) and v.strip().upper()=='TOTAL':
                if mcol is not None and mcol<len(row) and isinstance(row[mcol],(int,float)):
                    total=float(row[mcol])
    s,e=_dates(period)
    return dict(period=period, start=s, end=e, total_mandays=total, _source=path.split('/')[-1])

if __name__=='__main__':
    import sys,json; print(json.dumps(parse_mandays(sys.argv[1]),indent=1))
